import os
import re
import gc
import uuid
import tempfile
import logging
import copy
import threading
from collections import OrderedDict
from pathlib import Path
from typing import List, Dict, Any, Optional

from src.config import settings
from src.doc_errors import DocumentReadError
from src.legacy_formats import extract_doc_paragraphs, extract_odt_paragraphs
from src.ocr import MIN_PAGE_TEXT_CHARS, ocr_pdf_pages, tesseract_available

logger = logging.getLogger(__name__)

# Microsoft Word COM automation is documented by Microsoft as unsupported for
# concurrent/server-side use and is not thread-safe - serialize all access to a
# single Word instance at a time to avoid RPC/COM apartment conflicts.
_word_automation_lock = threading.Lock()

# Supported file extensions
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md", ".rtf", ".odt"}


def _format_span_text(span: dict) -> str:
    """Wraps a PyMuPDF text span in lightweight markdown (**bold**, *italic*)
    based on its font flags, so emphasis in the source document isn't
    silently lost on the AI reviewer."""
    text = span.get("text", "")
    if not text.strip():
        return text
    flags = span.get("flags", 0)
    font_name = (span.get("font") or "").lower()
    is_bold = bool(flags & (1 << 4)) or "bold" in font_name
    is_italic = bool(flags & (1 << 1)) or "italic" in font_name or "oblique" in font_name
    if is_bold and is_italic:
        return f"***{text}***"
    if is_bold:
        return f"**{text}**"
    if is_italic:
        return f"*{text}*"
    return text


def _format_dict_block(block: dict) -> str:
    """Reconstructs a PyMuPDF dict-mode text block into a string, preserving
    bold/italic emphasis markers per span."""
    lines = []
    for line in block.get("lines", []):
        line_text = "".join(_format_span_text(span) for span in line.get("spans", []))
        lines.append(line_text)
    return "\n".join(lines)


def _table_to_markdown(rows: list) -> str:
    """Converts a PyMuPDF extracted table (list of row lists) into a simple
    markdown table so row/column relationships survive as plain text instead
    of collapsing into disordered flowing paragraphs."""
    cleaned_rows = [
        [("" if cell is None else str(cell)).strip().replace("\n", " ") for cell in row]
        for row in rows
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    if not cleaned_rows:
        return ""

    lines = ["| " + " | ".join(cleaned_rows[0]) + " |"]
    lines.append("| " + " | ".join("---" for _ in cleaned_rows[0]) + " |")
    for row in cleaned_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


# Header/footer noise: short blocks near the page margin that are just a
# page number (e.g. "- 3 -", "Halaman 12") rather than real content.
_PAGE_NUMBER_NOISE_RE = re.compile(
    r"^[\-–—\s]*(halaman\s*)?\d{1,4}[\-–—\s]*$", re.IGNORECASE
)


def _extract_pdf(path: Path, allow_ocr: bool = True) -> List[Dict[str, Any]]:
    """
    Extract text from PDF page-by-page. Preserves bold/italic emphasis as
    lightweight markdown, extracts tables as markdown tables (instead of
    letting them collapse into disordered paragraph text), and filters out
    header/footer page-number noise near the page margins.
    """
    import pymupdf as fitz

    doc = fitz.open(str(path))
    pages: List[Dict[str, Any]] = []

    for page_idx in range(len(doc)):
        page = doc[page_idx]
        page_height = page.rect.height
        margin = page_height * 0.1

        # Detect tables first so their area can be excluded from normal
        # paragraph extraction and emitted as clean markdown instead.
        table_rects = []
        table_entries = []  # (x0, y0, x1, y1, markdown_text)
        try:
            found_tables = page.find_tables()
            for tbl in found_tables.tables:
                rect = fitz.Rect(tbl.bbox)
                table_rects.append(rect)
                md = _table_to_markdown(tbl.extract())
                if md:
                    table_entries.append((rect.x0, rect.y0, rect.x1, rect.y1, md))
        except Exception as e:
            logger.debug(f"Deteksi tabel dilewati di halaman {page_idx + 1}: {e}")

        raw_dict = page.get_text("dict")
        text_entries = []  # (x0, y0, x1, y1, text)
        for block in raw_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            bbox = block.get("bbox", (0, 0, 0, 0))
            block_rect = fitz.Rect(bbox)
            if any(block_rect.intersects(tr) for tr in table_rects):
                continue  # sudah ditangani sebagai tabel

            formatted = _format_dict_block(block).strip()
            if not formatted:
                continue

            is_near_margin = bbox[1] < margin or bbox[1] > (page_height - margin)
            is_noise = bool(_PAGE_NUMBER_NOISE_RE.match(formatted)) or len(formatted) <= 2
            if is_near_margin and is_noise:
                continue

            text_entries.append((bbox[0], bbox[1], bbox[2], bbox[3], formatted))

        combined = text_entries + table_entries
        combined.sort(key=lambda b: (b[1], b[0]))

        full_text = "\n\n".join(_merge_wrapped_lines_into_paragraphs(combined))
        pages.append({"page_number": page_idx + 1, "text": full_text})

    # Halaman tanpa teks (hasil scan) dibaca dengan OCR; halaman bergambar yang sudah punya teks dibiarkan.
    if allow_ocr:
        _apply_ocr(doc, pages)
    doc.close()
    return pages


def _merge_wrapped_lines_into_paragraphs(raw_blocks: list) -> List[str]:
    """
    Some PDFs (notably ones exported by Microsoft Word) have PyMuPDF report one
    "block" per visual LINE instead of per paragraph, which used to make every
    wrapped line look like its own paragraph once joined with blank lines.
    This reconstructs real paragraphs by merging consecutive blocks whose
    vertical gap matches normal line spacing, and only starts a new paragraph
    when the gap is notably larger (an actual blank line on the page).

    Args:
        raw_blocks: PyMuPDF text blocks (tuples), already filtered to text
            blocks and sorted top-to-bottom, left-to-right.
    """
    if not raw_blocks:
        return []

    gaps = [
        raw_blocks[i][1] - raw_blocks[i - 1][3]
        for i in range(1, len(raw_blocks))
    ]
    positive_gaps = [g for g in gaps if g > 0]
    normal_gap = sorted(positive_gaps)[len(positive_gaps) // 2] if positive_gaps else 0
    paragraph_gap_threshold = (normal_gap * 1.8) if normal_gap > 0 else 0

    paragraphs: List[str] = []
    current_lines: List[str] = []
    prev_y1 = None

    def flush_current():
        if current_lines:
            paragraphs.append(" ".join(current_lines))
            current_lines.clear()

    for b in raw_blocks:
        raw_text = b[4]
        # Pre-formatted markdown tables (see _table_to_markdown) must stay
        # intact as their own paragraph rather than being flattened with " ".
        # A plain text block spanning multiple internal lines is NOT atomic -
        # its line breaks are just normal word-wrap and should collapse to
        # spaces like everything else.
        is_atomic = raw_text.strip().startswith("| ")

        is_new_paragraph = (
            prev_y1 is not None
            and paragraph_gap_threshold > 0
            and (b[1] - prev_y1) > paragraph_gap_threshold
        )
        if is_new_paragraph or is_atomic:
            flush_current()

        if is_atomic:
            paragraphs.append(raw_text.strip())
        else:
            current_lines.append(" ".join(raw_text.split()))

        prev_y1 = b[3]

    flush_current()

    return paragraphs


def _extract_docx(path: Path) -> List[Dict[str, Any]]:
    """Extract text from DOCX document."""
    from docx import Document

    doc = Document(str(path))
    paragraphs: List[str] = []

    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paragraphs.append(t)

    # Also extract tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paragraphs.append(row_text)

    # Approximate pagination (roughly ~350 words per page)
    pages: List[Dict[str, Any]] = []
    current_page_paras: List[str] = []
    current_word_count = 0
    page_num = 1

    for p in paragraphs:
        w_count = len(p.split())
        if current_word_count + w_count > 350 and current_page_paras:
            pages.append({
                "page_number": page_num,
                "text": "\n\n".join(current_page_paras),
            })
            page_num += 1
            current_page_paras = [p]
            current_word_count = w_count
        else:
            current_page_paras.append(p)
            current_word_count += w_count

    if current_page_paras:
        pages.append({
            "page_number": page_num,
            "text": "\n\n".join(current_page_paras),
        })

    if not pages:
        pages.append({"page_number": 1, "text": ""})

    return pages


def _extract_rtf(path: Path) -> List[Dict[str, Any]]:
    """Extract plain text from RTF file."""
    from striprtf.striprtf import rtf_to_text

    raw = path.read_text(encoding="utf-8", errors="ignore")
    plain = rtf_to_text(raw)
    return _split_text_to_pages(plain)


def _extract_plain_text(path: Path) -> List[Dict[str, Any]]:
    """Extract text from TXT or Markdown file."""
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            content = path.read_text(encoding=enc)
            return _split_text_to_pages(content)
        except UnicodeDecodeError:
            continue
    content = path.read_text(encoding="utf-8", errors="ignore")
    return _split_text_to_pages(content)


def _split_text_to_pages(text: str, words_per_page: int = 350) -> List[Dict[str, Any]]:
    """Helper to partition raw continuous text into numbered pages."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [p.strip() for p in text.split("\n") if p.strip()]

    pages: List[Dict[str, Any]] = []
    current_paras: List[str] = []
    current_words = 0
    page_num = 1

    for p in paragraphs:
        w_count = len(p.split())
        if current_words + w_count > words_per_page and current_paras:
            pages.append({
                "page_number": page_num,
                "text": "\n\n".join(current_paras),
            })
            page_num += 1
            current_paras = [p]
            current_words = w_count
        else:
            current_paras.append(p)
            current_words += w_count

    if current_paras:
        pages.append({
            "page_number": page_num,
            "text": "\n\n".join(current_paras),
        })

    if not pages:
        pages.append({"page_number": 1, "text": ""})

    return pages


def _convert_docx_to_pdf_via_word(path: Path) -> Optional[Path]:
    """
    Uses locally installed Microsoft Word (COM automation, Windows-only) to convert
    a DOCX/DOC file to PDF, preserving Word's own real pagination. This gives page
    numbers that exactly match what the user sees when opening the file in Word,
    unlike the word-count-based approximation used as a fallback.

    Returns the temp PDF path on success, or None if Word automation is unavailable
    or fails for any reason (caller should fall back to the approximate extractor).
    """
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None

    pdf_path = Path(tempfile.gettempdir()) / f"{path.stem}_{uuid.uuid4().hex[:8]}.pdf"

    # Word automation is not safe for concurrent calls from multiple threads -
    # only one conversion runs at a time across the whole process.
    with _word_automation_lock:
        pythoncom.CoInitialize()
        word = None
        word_doc = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            word_doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)
            # wdFormatPDF = 17
            word_doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
            return pdf_path
        except Exception as e:
            logger.warning(
                f"Konversi DOCX->PDF via MS Word gagal untuk '{path.name}' ({e}), "
                f"pakai estimasi halaman berbasis jumlah kata."
            )
            return None
        finally:
            try:
                if word_doc is not None:
                    word_doc.Close(SaveChanges=0)
            except Exception:
                pass
            try:
                if word is not None:
                    word.Quit(SaveChanges=0)
            except Exception:
                pass
            # Release COM references explicitly before uninitializing - relying
            # on garbage collection alone can leave WINWORD.EXE running as an
            # orphaned process.
            word_doc = None
            word = None
            gc.collect()
            pythoncom.CoUninitialize()


# ---------------------------------------------------------------------------
# Titik masuk pembacaan dokumen
# ---------------------------------------------------------------------------
_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")  # penanda file Word 97-2003 (.doc)
MIN_TOTAL_TEXT_CHARS = 20

# Hasil pembacaan beberapa dokumen terakhir disimpan sementara: server membaca dokumen dua kali
# (sekali untuk menghitung blok, sekali saat pemeriksaan), dan OCR tidak boleh dikerjakan dua kali.
_CACHE_MAX = 3
_cache: "OrderedDict[tuple, List[Dict[str, Any]]]" = OrderedDict()
_cache_lock = threading.Lock()


def _looks_binary(sample: str) -> bool:
    """True bila teks didominasi karakter kendali / karakter pengganti (tanda file biner dibaca sebagai teks)."""
    if not sample:
        return False
    bad = sum(1 for ch in sample if (ord(ch) < 32 and ch not in "\n\r\t") or ch == "�")
    return bad / len(sample) > 0.02


def _validate_pages(pages: List[Dict[str, Any]], path: Path) -> None:
    """Tolak dokumen tanpa teks atau berisi data biner, supaya tidak diproses AI dan membuang kuota."""
    total = sum(len(p.get("text", "").strip()) for p in pages)
    if total < MIN_TOTAL_TEXT_CHARS:
        if path.suffix.lower() == ".pdf":
            raise DocumentReadError(
                "Tidak ada teks yang bisa dibaca dari PDF ini. Kemungkinan PDF hasil scan dengan gambar yang terlalu "
                "buram atau gelap sehingga OCR tidak mengenali tulisannya."
            )
        raise DocumentReadError("Tidak ada teks yang bisa dibaca dari dokumen ini.")
    sample = " ".join(p.get("text", "") for p in pages)[:20000]
    if _looks_binary(sample):
        raise DocumentReadError(
            "Isi file tampak berupa data biner, bukan teks dokumen. Pastikan formatnya sesuai ekstensi file, atau "
            "simpan ulang sebagai .docx atau .pdf."
        )


def _apply_ocr(doc: Any, pages: List[Dict[str, Any]]) -> None:
    """Baca dengan OCR halaman PDF yang tidak punya teks (hasil scan)."""
    need = [i for i, p in enumerate(pages) if len(p["text"].strip()) < MIN_PAGE_TEXT_CHARS]
    if not need:
        return
    has_any_text = any(len(p["text"].strip()) >= MIN_PAGE_TEXT_CHARS for p in pages)
    if has_any_text and (not settings.AI_OCR_ENABLED or not tesseract_available()):
        # Dokumen campuran dan OCR tidak ada: halaman bergambar dilewati, sisanya tetap diperiksa.
        logger.warning(f"{len(need)} halaman PDF tanpa teks dilewati karena OCR tidak tersedia.")
        return
    # Melempar DocumentReadError bila OCR tidak tersedia (dokumen seluruhnya gambar) atau terlalu panjang.
    ocr_pdf_pages(doc, pages, need)


def _extract_word(path: Path) -> List[Dict[str, Any]]:
    """Baca .docx / .doc. Mengenali isi sebenarnya (bukan hanya ekstensi): .doc yang ternyata .docx atau .rtf tetap terbaca."""
    # 1) Windows dengan MS Word terpasang: konversi ke PDF agar nomor halaman persis seperti di Word.
    pdf_path = _convert_docx_to_pdf_via_word(path)
    if pdf_path is not None:
        try:
            pages = _extract_pdf(pdf_path, allow_ocr=False)
            logger.info(
                f"'{path.name}' dikonversi via MS Word ke PDF sementara untuk nomor halaman yang akurat ({len(pages)} halaman)."
            )
            return pages
        except Exception as e:
            logger.warning(f"Gagal membaca hasil konversi PDF ({e}), pakai estimasi halaman.")
        finally:
            try:
                pdf_path.unlink(missing_ok=True)
            except Exception:
                pass

    with open(path, "rb") as f:
        head = f.read(8)
    if head.startswith(b"{\\rtf"):
        return _extract_rtf(path)
    if head.startswith(_OLE_MAGIC):
        return _split_text_to_pages("\n\n".join(extract_doc_paragraphs(path)))
    try:
        return _extract_docx(path)
    except Exception as e:
        logger.warning(f"Failed extracting DOCX via python-docx ({e}), trying text fallback.")
        return _extract_plain_text(path)


def _extract_uncached(path: Path, allow_ocr: bool) -> List[Dict[str, Any]]:
    ext = path.suffix.lower()
    logger.info(f"Extracting text from '{path.name}' (format: {ext})...")

    if ext == ".pdf":
        return _extract_pdf(path, allow_ocr=allow_ocr)
    if ext in (".docx", ".doc"):
        return _extract_word(path)
    if ext == ".rtf":
        return _extract_rtf(path)
    if ext == ".odt":
        return _split_text_to_pages("\n\n".join(extract_odt_paragraphs(path)))
    # .txt, .md, dan format lain berbasis teks
    return _extract_plain_text(path)


def extract_document_pages(file_path: str, allow_ocr: bool = True, strict: bool = True) -> List[Dict[str, Any]]:
    """
    Pembaca dokumen universal: PDF (termasuk hasil scan lewat OCR), DOCX, DOC, TXT, MD, RTF, dan ODT.

    Args:
        file_path: Path dokumen.
        allow_ocr: Baca halaman PDF tanpa teks dengan OCR. Dimatikan untuk indeks pedoman.
        strict: Tolak dokumen tanpa teks / berisi data biner dengan DocumentReadError.

    Returns:
        List of dicts: [{"page_number": int, "text": str, "ocr": True (hanya bila dibaca dengan OCR)}]
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File dokumen tidak ditemukan di: {file_path}")

    stat = path.stat()
    key = (str(path.resolve()), stat.st_mtime_ns, stat.st_size, allow_ocr, strict)
    with _cache_lock:
        if key in _cache:
            _cache.move_to_end(key)
            return copy.deepcopy(_cache[key])

    pages = _extract_uncached(path, allow_ocr)
    if strict:
        _validate_pages(pages, path)

    with _cache_lock:
        _cache[key] = pages
        while len(_cache) > _CACHE_MAX:
            _cache.popitem(last=False)
    return copy.deepcopy(pages)
