import io
import os
import re
import shutil
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from src.config import settings
from src.doc_errors import DocumentReadError
from src.document_extractor import _looks_binary, extract_document_pages
from src.ocr import tesseract_available

FIXTURES = Path(__file__).parent / "fixtures"

ISI = (
    "PERATURAN DAERAH PROVINSI RIAU\n"
    "NOMOR 1 TAHUN 2026 TENTANG PENYELENGGARAAN KETERTIBAN UMUM\n\n"
    "Pasal 1\n"
    "Dalam Peraturan Daerah ini yang dimaksud dengan Daerah adalah Provinsi Riau. Setiap orang wajib menjaga "
    "ketertiban umum dan menaati peraturan yang berlaku di wilayah Provinsi Riau tanpa terkecuali.\n\n"
    "Pasal 2\n"
    "Gubernur berwenang menetapkan peraturan pelaksanaan dari Peraturan Daerah ini sesuai dengan ketentuan "
    "peraturan perundang-undangan yang berlaku."
)

needs_tesseract = pytest.mark.skipif(not tesseract_available(), reason="tesseract tidak terpasang")
needs_antiword = pytest.mark.skipif(shutil.which("antiword") is None, reason="antiword tidak terpasang")


def _teks(pages):
    return " ".join(p["text"] for p in pages)


def _words(text):
    return re.findall(r"[a-z]+", text.lower())


def _tulis(tmp_path, nama, data):
    p = tmp_path / nama
    p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return str(p)


def _pdf_berteks():
    pdf = pymupdf.open()
    pdf.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 750), ISI, fontsize=11)
    return pdf.tobytes()


def _pdf_scan(jumlah_halaman=1):
    """PDF yang halamannya hanya berisi gambar (tanpa teks), seperti hasil scan."""
    sumber = pymupdf.open()
    sumber.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 750), ISI, fontsize=12)
    png = sumber[0].get_pixmap(dpi=200).tobytes("png")
    scan = pymupdf.open()
    for _ in range(jumlah_halaman):
        scan.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=png)
    return scan.tobytes()


# ---------------------------------------------------------------- format yang dulu sudah berfungsi
def test_txt_md_rtf_docx_dan_pdf_berteks_terbaca(tmp_path):
    doc = Document()
    for baris in ISI.split("\n"):
        if baris.strip():
            doc.add_paragraph(baris)
    buf = io.BytesIO()
    doc.save(buf)
    rtf = r"{\rtf1\ansi " + ISI.replace("\n", r"\par ") + "}"
    kasus = {
        "a.txt": ISI,
        "a.md": "# Peraturan\n\n" + ISI,
        "a.rtf": rtf,
        "a.docx": buf.getvalue(),
        "a.pdf": _pdf_berteks(),
    }
    for nama, isi in kasus.items():
        pages = extract_document_pages(_tulis(tmp_path, nama, isi))
        assert "Provinsi Riau" in _teks(pages), nama
        assert not any(p.get("ocr") for p in pages), f"{nama} berteks tidak perlu OCR"


# ---------------------------------------------------------------- .odt
def _odt_terkompresi():
    bo = io.BytesIO()
    with zipfile.ZipFile(bo, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text", compress_type=zipfile.ZIP_STORED)
        ps = "".join(f"<text:p>{x}</text:p>" for x in ISI.split("\n") if x.strip())
        z.writestr(
            "content.xml",
            '<?xml version="1.0"?><office:document-content '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            f"<office:body><office:text>{ps}</office:text></office:body></office:document-content>",
        )
    return bo.getvalue()


def test_odt_dibaca_sebagai_teks_bukan_data_biner(tmp_path):
    teks = _teks(extract_document_pages(_tulis(tmp_path, "a.odt", _odt_terkompresi())))
    assert "Provinsi Riau" in teks and "Pasal 2" in teks
    assert "PK" not in teks[:10] and "\x00" not in teks


def test_odt_rusak_ditolak_dengan_pesan_jelas(tmp_path):
    with pytest.raises(DocumentReadError, match="odt"):
        extract_document_pages(_tulis(tmp_path, "rusak.odt", b"bukan zip sama sekali"))


# ---------------------------------------------------------------- .doc
@needs_antiword
def test_doc_asli_buatan_word_terbaca_termasuk_tabel():
    pages = extract_document_pages(str(FIXTURES / "sample.doc"))
    teks = _teks(pages)
    assert "PERATURAN DAERAH PROVINSI RIAU" in teks
    assert "Pasal 1" in teks and "Pasal 2" in teks
    assert "waijb" in teks  # kesalahan ketik sengaja di berkas uji harus tetap terbaca apa adanya
    assert "Ketertiban umum" in teks  # isi sel tabel


def test_doc_yang_sebenarnya_docx_atau_rtf_tetap_terbaca(tmp_path):
    doc = Document()
    doc.add_paragraph(ISI.replace("\n", " "))
    buf = io.BytesIO()
    doc.save(buf)
    assert "Provinsi Riau" in _teks(extract_document_pages(_tulis(tmp_path, "salah_ekstensi.doc", buf.getvalue())))
    rtf = r"{\rtf1\ansi " + ISI.replace("\n", r"\par ") + "}"
    assert "Provinsi Riau" in _teks(extract_document_pages(_tulis(tmp_path, "rtf_di_doc.doc", rtf)))


def test_doc_acak_ditolak_bukan_dikirim_ke_ai(tmp_path):
    ole_palsu = bytes.fromhex("D0CF11E0A1B11AE1") + os.urandom(4096)
    with pytest.raises(DocumentReadError):
        extract_document_pages(_tulis(tmp_path, "acak.doc", ole_palsu))


# ---------------------------------------------------------------- penolakan dokumen tanpa teks
def test_dokumen_kosong_ditolak(tmp_path):
    with pytest.raises(DocumentReadError, match="Tidak ada teks"):
        extract_document_pages(_tulis(tmp_path, "kosong.txt", "   \n\n "))


def test_data_biner_yang_menyamar_sebagai_txt_ditolak(tmp_path):
    with pytest.raises(DocumentReadError, match="biner"):
        extract_document_pages(_tulis(tmp_path, "biner.txt", bytes(range(256)) * 20))


def test_deteksi_biner():
    assert _looks_binary("teks biasa dengan\tbaris\nbaru") is False
    assert _looks_binary("PK\x03\x04\x14\x00\x00\x00\x00\x00abc" * 10) is True
    assert _looks_binary("") is False


# ---------------------------------------------------------------- PDF hasil scan (OCR)
@needs_tesseract
def test_pdf_hasil_scan_dibaca_dengan_ocr(tmp_path):
    pages = extract_document_pages(_tulis(tmp_path, "scan.pdf", _pdf_scan()))
    assert pages and all(p.get("ocr") for p in pages), "halaman gambar harus ditandai hasil OCR"
    asli = set(_words(ISI))
    terbaca = _words(_teks(pages))
    cocok = sum(1 for w in terbaca if w in asli) / max(1, len(terbaca))
    assert cocok >= 0.85, f"hanya {cocok:.0%} kata hasil OCR cocok dengan teks asli"
    assert "provinsi" in terbaca and "riau" in terbaca and "pasal" in terbaca


@needs_tesseract
def test_hasil_ocr_disimpan_sementara_sehingga_tidak_diulang(tmp_path, monkeypatch):
    import src.document_extractor as de

    path = _tulis(tmp_path, "scan.pdf", _pdf_scan())
    panggilan = []
    asli = de.ocr_pdf_pages
    monkeypatch.setattr(de, "ocr_pdf_pages", lambda *a, **k: panggilan.append(1) or asli(*a, **k))
    extract_document_pages(path)
    extract_document_pages(path)
    assert len(panggilan) == 1


@needs_tesseract
def test_pdf_scan_terlalu_panjang_ditolak(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "AI_OCR_MAX_PAGES", 1)
    with pytest.raises(DocumentReadError, match="terlalu panjang"):
        extract_document_pages(_tulis(tmp_path, "scan2.pdf", _pdf_scan(jumlah_halaman=2)))


def test_pdf_scan_saat_ocr_dimatikan_ditolak_dengan_pesan_jelas(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "AI_OCR_ENABLED", False)
    with pytest.raises(DocumentReadError, match="OCR tidak tersedia"):
        extract_document_pages(_tulis(tmp_path, "scan3.pdf", _pdf_scan()))


def test_jalur_pedoman_tidak_memakai_ocr_dan_tidak_melempar_error(tmp_path):
    # index_pedoman memanggil allow_ocr=False, strict=False: PDF gambar menghasilkan halaman kosong, bukan error.
    pages = extract_document_pages(_tulis(tmp_path, "pedoman_scan.pdf", _pdf_scan()), allow_ocr=False, strict=False)
    assert len(pages) == 1 and pages[0]["text"].strip() == "" and not pages[0].get("ocr")
