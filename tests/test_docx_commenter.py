import io

from docx import Document
from docx.shared import Cm, Pt

from src.docx_commenter import annotate_docx
from src.schemas import BlockReviewResult, CheckReport, RuleReference, TextSpanError


def _make_docx() -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(3.5), Cm(2.0)

    # Kata "waijb" terbelah di dua run berformat berbeda (run kedua tebal).
    p1 = doc.add_paragraph()
    p1.add_run("Setiap orang wa")
    r2 = p1.add_run("ijb menjaga ketertiban umum.")
    r2.bold = True
    r2.font.size = Pt(13)

    # Spasi ganda di dokumen, sedangkan teks blok dari ekstraktor sudah ternormalisasi.
    doc.add_paragraph("Daerah  adalah   Provinsi Riau.")

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nomer"
    table.cell(0, 1).text = "Keterangan"
    return _to_bytes(doc)


def _to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _err(snippet, fix, tipe="ejaan"):
    return TextSpanError(
        original_snippet=snippet,
        suggested_snippet=fix,
        error_type=tipe,
        explanation=f"'{snippet}' tidak baku",
        severity="high",
    )


def _report(blocks):
    return CheckReport(document_checked="uji.docx", summary={"total_blocks": len(blocks)}, blocks=blocks)


def _annotate(report):
    out, stats = annotate_docx(_make_docx(), report)
    return Document(io.BytesIO(out.getvalue())), stats


def _comment_texts(doc):
    return {c.comment_id: c.text for c in doc.comments}


def _anchored_text(doc, comment_id):
    """Teks yang dicakup rentang komentar (antara commentRangeStart dan commentRangeEnd)."""
    from docx.oxml.ns import qn

    body = doc.element.body
    start = body.xpath(f".//w:commentRangeStart[@w:id='{comment_id}']")[0]
    out, node, inside = [], start, True
    node = start.getnext()
    while node is not None and node.tag != qn("w:commentRangeEnd"):
        out.extend(t.text or "" for t in node.iter(qn("w:t")))
        node = node.getnext()
    return "".join(out)


def test_komentar_menempel_pada_kata_yang_salah_walau_terbelah_run():
    block = BlockReviewResult(
        block_id="block_1",
        original_text="Setiap orang waijb menjaga ketertiban umum.",
        status="ejaan_tanda_baca",
        span_errors=[_err("waijb", "wajib")],
    )
    doc, stats = _annotate(_report([block]))
    assert stats.placed == 1 and stats.unplaced == 0
    comments = list(doc.comments)
    assert len(comments) == 1
    assert "wajib" in comments[0].text
    assert _anchored_text(doc, comments[0].comment_id) == "waijb"


def test_isi_dan_format_dokumen_tidak_berubah():
    before = Document(io.BytesIO(_make_docx()))
    block = BlockReviewResult(
        block_id="block_1",
        original_text="Setiap orang waijb menjaga ketertiban umum.",
        status="ejaan_tanda_baca",
        span_errors=[_err("waijb", "wajib")],
    )
    after, _ = _annotate(_report([block]))

    assert [p.text for p in after.paragraphs] == [p.text for p in before.paragraphs]
    assert after.tables[0].cell(0, 0).text == "Nomer"

    s0, s1 = before.sections[0], after.sections[0]
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin", "page_width", "page_height"):
        assert getattr(s0, attr) == getattr(s1, attr)

    # Format run kedua (tebal, 13pt) harus tetap pada seluruh potongannya.
    runs = after.paragraphs[0].runs
    bold_text = "".join(r.text for r in runs if r.bold)
    assert bold_text == "ijb menjaga ketertiban umum."
    assert all(r.font.size == Pt(13) for r in runs if r.bold)


def test_pencocokan_mengabaikan_spasi_ganda():
    block = BlockReviewResult(
        block_id="block_2",
        original_text="Daerah adalah Provinsi Riau.",
        status="ejaan_tanda_baca",
        span_errors=[_err("Daerah adalah", "Daerah ialah", "kosa_kata")],
    )
    doc, stats = _annotate(_report([block]))
    assert stats.placed == 1
    cid = list(doc.comments)[0].comment_id
    assert _anchored_text(doc, cid) == "Daerah  adalah"  # teks asli dengan spasi ganda tetap utuh


def test_kata_di_dalam_tabel_juga_dikomentari():
    block = BlockReviewResult(
        block_id="block_3",
        original_text="Nomer | Keterangan",
        status="ejaan_tanda_baca",
        span_errors=[_err("Nomer", "Nomor", "kosa_kata")],
    )
    doc, stats = _annotate(_report([block]))
    assert stats.placed == 1
    cid = list(doc.comments)[0].comment_id
    assert _anchored_text(doc, cid) == "Nomer"


def test_blok_perlu_revisi_diberi_komentar_tingkat_blok_beserta_rujukan():
    block = BlockReviewResult(
        block_id="block_1",
        original_text="Setiap orang waijb menjaga ketertiban umum.",
        status="perlu_revisi",
        issue="Rumusan norma bertentangan dengan UU 12/2011.",
        suggested_revision="Setiap orang wajib menjaga ketertiban umum.",
        rule_reference=RuleReference(document="UU 12 Tahun 2011.pdf", page=17),
        span_errors=[],
    )
    doc, stats = _annotate(_report([block]))
    assert stats.placed == 1
    text = list(doc.comments)[0].text
    assert "bertentangan" in text and "UU 12 Tahun 2011.pdf" in text and "Usulan revisi" in text


def test_kemunculan_berulang_diberi_komentar_sendiri_sendiri():
    doc0 = Document()
    doc0.add_paragraph("waijb dan waijb lagi.")
    buf = io.BytesIO()
    doc0.save(buf)
    block = BlockReviewResult(
        block_id="block_1",
        original_text="waijb dan waijb lagi.",
        status="ejaan_tanda_baca",
        span_errors=[_err("waijb", "wajib"), _err("waijb", "wajib")],
    )
    out, stats = annotate_docx(buf.getvalue(), _report([block]))
    doc = Document(io.BytesIO(out.getvalue()))
    assert stats.placed == 2 and len(list(doc.comments)) == 2
    texts = [_anchored_text(doc, c.comment_id) for c in doc.comments]
    assert texts == ["waijb", "waijb"]
    assert doc.paragraphs[0].text == "waijb dan waijb lagi."


def test_snippet_yang_tidak_ada_dilaporkan_sebagai_tidak_tertempel():
    block = BlockReviewResult(
        block_id="block_1",
        original_text="Setiap orang waijb menjaga ketertiban umum.",
        status="ejaan_tanda_baca",
        span_errors=[_err("kata_yang_tidak_ada", "x")],
    )
    doc, stats = _annotate(_report([block]))
    assert stats.placed == 0 and stats.unplaced == 1
    assert len(list(doc.comments)) == 0
